import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import re
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_styled_document(md_path, docx_path):
    doc = docx.Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        # Parse table markdown
        rows_data = []
        for line in t_lines:
            line = line.strip()
            if not line or line.startswith('|---') or line.startswith('|:---'):
                continue
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if cells:
                rows_data.append(cells)
        if not rows_data:
            return

        cols_cnt = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=cols_cnt)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for i, row in enumerate(rows_data):
            for j, cell_text in enumerate(row):
                if j < cols_cnt:
                    cell = table.cell(i, j)
                    cell.text = cell_text
                    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.line_spacing = 1.1
                    
                    if i == 0: # Header
                        set_cell_background(cell, "1B365D")
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else: # Body
                        bg = "F0F4F8" if i % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg)
                        for run in p.runs:
                            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        doc.add_paragraph() # Spacing after table

    def format_runs(p, text):
        # Support bold **text** and inline code `code`
        parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('`') and part.endswith('`'):
                run = p.add_run(part[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
            else:
                p.add_run(part)

    for line in lines:
        raw_line = line.rstrip('\n\r')

        # Handle code blocks
        if raw_line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                # Write code block paragraph
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.right_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)

                # Background box setup
                pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="12" w:color="1B365D"/></w:pBdr>')
                p._element.get_or_add_pPr().append(pBdr)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F9"/>')
                p._element.get_or_add_pPr().append(shd)

                code_text = "\n".join(code_lines)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
                code_lines = []
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(raw_line)
            continue

        # Handle tables
        if raw_line.strip().startswith('|'):
            if not in_table:
                in_table = True
            table_lines.append(raw_line)
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        stripped = raw_line.strip()

        if not stripped:
            continue

        # Headings
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(stripped[2:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x54, 0x8A)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[4:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x3B, 0x6E, 0xA5)
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[5:])
            run.font.name = 'Segoe UI'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="3B6EA5"/></w:pBdr>')
            p._element.get_or_add_pPr().append(pBdr)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
            p._element.get_or_add_pPr().append(shd)
            format_runs(p, stripped[2:])
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            format_runs(p, stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            sub_text = re.sub(r'^\d+\.\s', '', stripped)
            format_runs(p, sub_text)
        elif stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            p._element.get_or_add_pPr().append(pBdr)
        else:
            p = doc.add_paragraph()
            format_runs(p, stripped)

    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print("Successfully saved docx file.")

if __name__ == '__main__':
    md = r"d:\Nexis\Training\System Design cho hệ thống PP Hub\README.md"
    out_docx = r"d:\Nexis\Training\System Design cho hệ thống PP Hub\Thiet_Ke_He_Thong_PP_Hub.docx"
    create_styled_document(md, out_docx)
